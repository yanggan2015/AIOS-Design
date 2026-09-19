# -*- coding: utf-8 -*-
"""Dense portrait slides: problem, mechanism, architecture. No slogan filler."""
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu, Inches

W, H = 1080, 1920
MX = 56
CW = W - MX * 2
BG_TOP = (22, 24, 28)
BG_BOT = (11, 12, 14)
CARD = (30, 33, 38)
LINE = (64, 58, 48)
GOLD = (216, 176, 106)
TEXT = (246, 243, 236)
MUTED = (176, 170, 160)
FONT_B = "C:/Windows/Fonts/msyhbd.ttc"
FONT_R = "C:/Windows/Fonts/msyh.ttc"

LINES = [
    "luminOS 是一套面向 AI Agent 的 Linux 操作系统。差的不是再做一个聊天框，靠的是操作系统自己提供桌面操控平面，应用不用预留 AI 接口也能被操作。",
    "现在的问题有三个。没有接口的图形软件，智能体进不去。每个应用单独做插件，覆盖永远接不完。坐标脚本没有目标绑定、没有结果检查，也没有权限和审计。",
    "解法一一对应。进不去，就用语义、协议、视觉三条通道，没有控件树也能在窗口内操作。接不完，是因为能力做在系统服务里，做一次所有应用都能用。脚本又脆，就改成每一步先快照、再动作、再断言，高风险必须确认。",
    "架构就四层。入口接收目标和远程报文。智能体负责理解和计划。操控平面负责找窗口并执行。最下面仍是普通 Linux 桌面，包括 Wayland、无障碍树和输入法。人工智能引擎只做推理，不进内核。",
    "三条通道按可靠度自动选择。能读到控件，就按角色和名称点击输入。系统已有消息或命令，就走协议。两者都没有，就截取这个窗口，识别文字，再在窗口内部注入，中文经输入法提交。",
    "执行不是盲点。动作绑定到具体窗口，带一次性令牌。做完要检查焦点、文本或界面是否变成预期。检查失败就停止，不用下一步去猜。密码、支付和认证界面默认拒绝自动化。",
    "权限分四级。读公开状态可以在会话内授权。输入和切换窗口用任务级信任。截图、写文件、装包要显式确认。发送、删除、重启和支付必须单次确认，并且全程可审计。",
    "网络和串口用同一套报文。帧头固定，后面是命令号、序号、载荷和校验。可以控制系统白名单功能，也可以发现窗口、点击和输入，还可以下发智能体任务。没有认证，这些命令一律拒绝。",
    "它承诺的是可安装、可解释、可收回。标准桌面应用走语义通道，没有接口的走视觉兜底。不承诺自动破解验证码，不承诺把大模型放进内核。你要的是万能像素脚本，还是一套带权限的系统操控平面？",
]


def font(path, size):
    return ImageFont.truetype(path, size)


def wrap(draw, text, fnt, max_w):
    lines, cur = [], ""
    for ch in text:
        trial = cur + ch
        if draw.textlength(trial, font=fnt) <= max_w:
            cur = trial
        else:
            if cur:
                lines.append(cur)
            cur = ch
    if cur:
        lines.append(cur)
    return lines


def blit(d, lines, x, y, fnt, fill, size, leading):
    for i, line in enumerate(lines):
        d.text((x, y + i * (size + leading)), line, font=fnt, fill=fill)
    return y + len(lines) * (size + leading)


def canvas():
    col = Image.new("RGB", (1, H))
    px = col.load()
    for y in range(H):
        t = y / (H - 1)
        px[0, y] = tuple(int(BG_TOP[i] + (BG_BOT[i] - BG_TOP[i]) * t) for i in range(3))
    im = col.resize((W, H))
    d = ImageDraw.Draw(im)
    d.rectangle((0, 0, W, 6), fill=GOLD)
    return im, d


def chrome(d, n, total, section):
    d.text((MX, 28), section, font=font(FONT_B, 22), fill=GOLD)
    label = f"{n:02d} / {total:02d}"
    f = font(FONT_R, 20)
    d.text((W - MX - d.textlength(label, font=f), 30), label, font=f, fill=MUTED)


def heading(d, text, y=78, size=42):
    f = font(FONT_B, size)
    return blit(d, wrap(d, text, f, CW), MX, y, f, TEXT, size, 8)


def card(d, y, h):
    d.rounded_rectangle((MX, y, MX + CW, y + h), radius=16, fill=CARD, outline=LINE, width=2)
    d.rectangle((MX, y + 16, MX + 5, y + h - 16), fill=GOLD)


def block(d, y, h, kicker, title, body):
    card(d, y, h)
    d.text((MX + 28, y + 18), kicker, font=font(FONT_B, 22), fill=GOLD)
    d.text((MX + 28, y + 52), title, font=font(FONT_B, 32), fill=TEXT)
    f = font(FONT_R, 26)
    blit(d, wrap(d, body, f, CW - 56), MX + 28, y + 108, f, MUTED, 26, 8)
    return y + h + 16


def main():
    stills = Path(r"C:\code\剪映skills\output\luminos_os\stills")
    stills.mkdir(parents=True, exist_ok=True)
    for old in stills.glob("*.png"):
        old.unlink()
    ppt_dir = Path(r"C:\code\ai操作系统的设计\宣传\竖屏")
    ppt_dir.mkdir(parents=True, exist_ok=True)
    script_path = Path(r"C:\code\剪映skills\output\luminos_os\script.txt")
    script_path.write_text("\n".join(LINES) + "\n", encoding="utf-8")

    slides = []

    def add(section, painter):
        slides.append((section, painter))

    def s1(d):
        d.text((MX, 120), "luminOS", font=font(FONT_B, 78), fill=TEXT)
        y = heading(d, "面向 AI Agent 的 Linux 操作系统", 230, 40)
        items = [
            ("操控面在系统", "应用不预留 AI 接口，也可被 Agent 操作。"),
            ("自带执行链", "Agent 规划、工具调用、结果检查都在本机。"),
            ("远程同一协议", "网络和串口共用报文，不另开一套权限。"),
        ]
        y = 430
        for t, b in items:
            y = block(d, y, 250, "定义", t, b)

    def s2(d):
        y = heading(d, "要解决的三个问题", 90, 44)
        rows = [
            ("01  进不去", "图形软件通常没有 Agent API。窗口在，智能体仍不能点、不能输入。"),
            ("02  接不完", "逐个应用写插件，覆盖速度取决于厂商，做不成操作系统能力。"),
            ("03  不可审计", "坐标宏没有窗口绑定、没有结果断言、没有分级确认，改版即失效。"),
        ]
        y += 20
        for t, b in rows:
            y = block(d, y, 300, "问题", t, b)

    def s3(d):
        y = heading(d, "每个问题对应一个机制", 90, 42)
        rows = [
            ("进不去 → 三通道", "语义树优先。有系统协议就调用。都没有，则窗口内截图、识别、注入。"),
            ("接不完 → 系统服务", "操控平面只实现一次。已安装应用默认零改造即可被调用。"),
            ("不可审计 → 事务", "前快照、一次性令牌、动作、断言。失败即停。高风险动作必须确认。"),
        ]
        y += 16
        for t, b in rows:
            y = block(d, y, 310, "解法", t, b)

    def s4(d):
        y = heading(d, "架构：四层，推理不进内核", 90, 40)
        layers = [
            ("01", "入口", "快捷键、命令行、MCP、网络与串口报文。"),
            ("02", "智能体", "意图、计划、预算、确认、熔断、审计。"),
            ("03", "操控平面", "发现目标、选择通道、执行、检查结果。"),
            ("04", "Linux 会话", "Wayland、无障碍树、输入法、门户授权。"),
        ]
        y += 12
        for num, name, desc in layers:
            card(d, y, 200)
            d.text((MX + 28, y + 58), num, font=font(FONT_B, 28), fill=GOLD)
            d.text((MX + 120, y + 28), name, font=font(FONT_B, 34), fill=TEXT)
            f = font(FONT_R, 26)
            blit(d, wrap(d, desc, f, CW - 160), MX + 120, y + 86, f, MUTED, 26, 8)
            y += 216
        f = font(FONT_R, 26)
        blit(d, wrap(d, "AI 引擎只负责路由和推理，权重与算子都在用户态。", f, CW), MX, y + 8, f, MUTED, 26, 8)

    def s5(d):
        y = heading(d, "操控平面的三条通道", 90, 42)
        rows = [
            ("A  语义", "读 AT-SPI 控件树。按角色、名称、状态执行点击和输入。抗分辨率变化。"),
            ("B  协议", "调用已有 D-Bus、系统 Intent 或命令行。不要求应用另做 AI 接口。"),
            ("C  合成", "只截目标窗口。OCR 定位。指针键盘注入窗口内部。中文走输入法提交。"),
        ]
        y += 12
        for t, b in rows:
            y = block(d, y, 300, "通道", t, b)
        f = font(FONT_R, 26)
        blit(d, ["选择顺序：语义，协议，合成。失败要给出原因码。"], MX, y + 4, f, MUTED, 26, 8)

    def s6(d):
        y = heading(d, "一步动作怎么落地", 90, 44)
        rows = [
            ("绑定窗口", "禁止无目标的全局乱点。坐标相对窗口客户区，记录缩放。"),
            ("令牌", "动作绑定任务、目标和参数摘要，短时有效，可立即吊销。"),
            ("断言", "做完检查焦点、文本或控件树是否达到预期。达不到就停止。"),
            ("拒绝区", "密码框、支付页、系统认证界面默认不做自动化。"),
        ]
        y += 12
        for t, b in rows:
            y = block(d, y, 240, "执行", t, b)

    def s7(d):
        y = heading(d, "权限四级，不是统一放行", 90, 40)
        rows = [
            ("L0", "读公开窗口状态", "会话内授权"),
            ("L1", "输入、切换窗口", "任务级信任"),
            ("L2", "截图、写文件、装包", "显式确认"),
            ("L3", "发送、删除、重启、支付", "单次确认 + 审计"),
        ]
        y += 16
        for lv, name, rule in rows:
            card(d, y, 180)
            d.text((MX + 28, y + 64), lv, font=font(FONT_B, 32), fill=GOLD)
            d.text((MX + 140, y + 36), name, font=font(FONT_B, 32), fill=TEXT)
            d.text((MX + 140, y + 92), rule, font=font(FONT_R, 26), fill=MUTED)
            y += 196
        f = font(FONT_R, 26)
        blit(d, wrap(d, "验证码、DRM、反自动化不绕过。模型看不见明文凭据。", f, CW), MX, y + 6, f, MUTED, 26, 8)

    def s8(d):
        y = heading(d, "网络与串口：同一套报文", 90, 40)
        rows = [
            ("帧", "固定帧头，命令号，序号，载荷长度，CRC。两种传输字节级一致。"),
            ("系统", "白名单方法：音量、锁屏、休眠、重启。没有任意命令行。"),
            ("桌面", "列出窗口、快照、查找、点击、输入法提交。无图形会话则拒绝。"),
            ("任务", "下发目标、回进度、中止注入。未认证连接不能发这些命令。"),
        ]
        y += 12
        for t, b in rows:
            y = block(d, y, 230, "RCP", t, b)

    def s9(d):
        y = heading(d, "承诺什么，不承诺什么", 90, 42)
        rows = [
            ("承诺", "基线内 Linux 应用可安装。标准控件走语义通道。无接口窗口仍可尝试合成通道。失败说明原因。"),
            ("承诺", "高风险动作可确认、可审计、可停止。远程与本机走同一套权限。"),
            ("不承诺", "不自动破解验证码和 DRM。不把模型放进内核。不保证任意像素任务成功。"),
        ]
        y += 16
        for t, b in rows:
            y = block(d, y, 310, "边界", t, b)

    for fn in (s1, s2, s3, s4, s5, s6, s7, s8, s9):
        add("", fn)
    sections = ["定位", "问题", "解法", "架构", "通道", "执行", "权限", "报文", "边界"]

    prs = Presentation()
    prs.slide_width = Inches(7.5)
    prs.slide_height = Inches(13.333333)
    blank = prs.slide_layouts[6]
    total = len(sections)
    for i, (section, fn) in enumerate(zip(sections, [s1, s2, s3, s4, s5, s6, s7, s8, s9]), 1):
        im, d = canvas()
        fn(d)
        chrome(d, i, total, section)
        path = stills / f"{i:02d}.png"
        im.save(path, "PNG")
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(path), Emu(0), Emu(0), prs.slide_width, prs.slide_height)
        slide.notes_slide.notes_text_frame.text = LINES[i - 1]
    ppt = ppt_dir / "luminOS宣传-竖屏.pptx"
    prs.save(ppt)
    print(ppt)


if __name__ == "__main__":
    main()

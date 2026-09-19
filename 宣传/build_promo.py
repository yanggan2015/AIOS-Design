# -*- coding: utf-8 -*-
"""Generate AIOS promo slides, PPTX, and MP4."""
import asyncio
import subprocess
from pathlib import Path

import edge_tts
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu, Inches

ROOT = Path(__file__).resolve().parent
SLIDES = ROOT / "slides"
AUDIO = ROOT / "audio"
SLIDES.mkdir(exist_ok=True)
AUDIO.mkdir(exist_ok=True)

W, H = 1920, 1080
BG = (10, 16, 28)
CARD = (20, 30, 46)
LINE = (42, 58, 82)
ACCENT = (62, 224, 176)
BLUE = (110, 168, 255)
TEXT = (241, 245, 249)
MUTED = (156, 170, 188)
AMBER = (251, 191, 36)

FONT_B = "C:/Windows/Fonts/msyhbd.ttc"
FONT_R = "C:/Windows/Fonts/msyh.ttc"
VOICE = "zh-CN-YunxiNeural"


def font(path, size):
    return ImageFont.truetype(path, size)


def wrap(draw, text, fnt, max_w):
    lines, cur = [], ""
    for ch in text:
        trial = cur + ch
        if draw.textlength(trial, font=fnt) <= max_w:
            cur = trial
        else:
            if cur:
                lines.append(cur)
            cur = ch
    if cur:
        lines.append(cur)
    return lines


def rounded(draw, box, r, fill, outline=None):
    draw.rounded_rectangle(box, radius=r, fill=fill, outline=outline, width=2)


def base():
    im = Image.new("RGB", (W, H), BG)
    d = ImageDraw.Draw(im)
    d.rectangle((0, 0, 10, H), fill=ACCENT)
    d.rectangle((0, H - 6, W, H), fill=ACCENT)
    return im, d


def footer(d, n, total, label="AIOS  ·  基于 Linux 的 AI 操作系统"):
    f = font(FONT_R, 22)
    d.text((72, 1020), label, font=f, fill=MUTED)
    d.text((1680, 1020), f"{n:02d}  /  {total:02d}", font=f, fill=MUTED)


def kicker(d, text):
    f = font(FONT_B, 26)
    d.text((72, 64), text.upper() if False else text, font=f, fill=ACCENT)


def title(d, text, y=120, size=64, color=TEXT):
    f = font(FONT_B, size)
    lines = wrap(d, text, f, 1700)
    for i, line in enumerate(lines):
        d.text((72, y + i * (size + 16)), line, font=f, fill=color)
    return y + len(lines) * (size + 16)


def body(d, text, y, size=32, color=MUTED, width=1700):
    f = font(FONT_R, size)
    lines = wrap(d, text, f, width)
    for i, line in enumerate(lines):
        d.text((72, y + i * (size + 14)), line, font=f, fill=color)
    return y + len(lines) * (size + 14)


def card(d, x, y, w, h, head, text, accent=ACCENT):
    rounded(d, (x, y, x + w, y + h), 18, CARD, LINE)
    d.rectangle((x, y, x + 8, y + h), fill=accent)
    hf = font(FONT_B, 32)
    bf = font(FONT_R, 26)
    d.text((x + 36, y + 28), head, font=hf, fill=TEXT)
    lines = wrap(d, text, bf, w - 72)
    yy = y + 88
    for line in lines[:5]:
        d.text((x + 36, yy), line, font=bf, fill=MUTED)
        yy += 40


SLIDE_DEFS = [
    {
        "id": "01",
        "narration": "这是 AIOS。一套基于 Linux 的人工智能操作系统。它的目标很直接：安装在这套系统上的桌面应用，即使没有预留 AI 接口，也可以被系统里的智能体操控。",
        "draw": "cover",
    },
    {
        "id": "02",
        "narration": "今天的智能体，大多只能操作那些愿意开门的软件。图形应用没有接口，AI 就只能看着。如果每个应用各自做一套插件，那就永远做不成操作系统。",
        "draw": "problem",
    },
    {
        "id": "03",
        "narration": "AIOS 的答案，是由操作系统自己长出操作面。我们把它叫做桌面操控平面，简称 DCP。智能体不对接每个应用的私有协议，只调用这一层系统服务。",
        "draw": "dcp",
    },
    {
        "id": "04",
        "narration": "技术上采用三通道融合。能读控件树，就走无障碍语义。有系统协议，就走精确调用。两者都没有，就用窗口截图、文字识别和窗口内注入来完成。没有接口，也能进入最后一条通道。",
        "draw": "modes",
    },
    {
        "id": "05",
        "narration": "系统自带 AI 引擎和智能体。本地模型负责隐私和离线，复杂任务可以按策略上云。计划可以先看再执行，成功的流程可以存成技能，下次直接重跑。",
        "draw": "agent",
    },
    {
        "id": "06",
        "narration": "接入做到零改造。应用不用改代码，就能被操控。你也可以继续用自己的模型和外部智能体。执行、截图和权限，仍然留在这台机器上。",
        "draw": "access",
    },
    {
        "id": "07",
        "narration": "控制不止发生在屏幕前。网络和串口使用同一套报文协议。可以启动应用、操作窗口、下发任务，也可以控制电源和音量这类系统能力。远程通道不另开后门。",
        "draw": "rcp",
    },
    {
        "id": "08",
        "narration": "敢把控制权交给系统，是因为权限被分级。高风险动作必须确认，密码和支付界面默认拒绝自动化，每一步都可以审计，也可以立刻停下来。",
        "draw": "security",
    },
    {
        "id": "09",
        "narration": "这套方案建立在 Linux 已经成熟的能力上：无障碍树、Wayland、门户、输入法和资源隔离。大模型不进内核，也不靠破解应用的私有协议。所以它是可落地的，不是概念演示。",
        "draw": "feasible",
    },
    {
        "id": "10",
        "narration": "AIOS 要做的，不是再造一个聊天窗口，而是让操作系统本身成为 AI 的双手。不在应用里找接口，在系统里提供桌面操控平面。",
        "draw": "end",
    },
]


def draw_cover(d):
    kicker(d, "PROMOTIONAL  ·  DESIGN v1.2")
    f = font(FONT_B, 120)
    d.text((72, 220), "AIOS", font=f, fill=TEXT)
    title(d, "让所有桌面应用，都能被 AI 操控", y=390, size=56)
    body(
        d,
        "基于 Linux 内核的 AI 操作系统。系统自带智能体与 AI 引擎，并由操作系统提供操控能力。",
        560,
        34,
    )
    pills = ["系统级操控", "应用零改造", "网络 / 串口同一协议"]
    x = 72
    pf = font(FONT_B, 26)
    for p in pills:
        tw = d.textlength(p, font=pf)
        rounded(d, (x, 760, x + tw + 48, 824), 28, CARD, ACCENT)
        d.text((x + 24, 776), p, font=pf, fill=ACCENT)
        x += tw + 72


def draw_problem(d):
    kicker(d, "问题")
    title(d, "没有 AI 接口的软件，今天几乎控不了", y=130, size=52)
    items = [
        ("应用不开门", "图形软件没有给智能体预留接口，自动化就停在窗口外面。"),
        ("插件做不完", "每个软件各自对接，覆盖速度取决于厂商，成不了操作系统。"),
        ("脚本太脆", "纯坐标宏一改版就失效，而且没有权限、确认和审计。"),
    ]
    for i, (h, t) in enumerate(items):
        card(d, 72 + i * 600, 420, 560, 420, h, t, AMBER if i == 0 else ACCENT)


def draw_dcp(d):
    kicker(d, "产品亮点")
    title(d, "操作系统自己长出操作面", y=130, size=56)
    body(d, "桌面操控平面 DCP：智能体只面对系统，不面对每个应用的私有协议。", 280, 32)
    steps = ["发现窗口", "选择通道", "执行动作", "验证结果"]
    for i, s in enumerate(steps):
        x = 72 + i * 450
        rounded(d, (x, 480, x + 400, 680), 18, CARD, LINE)
        nf = font(FONT_B, 42)
        sf = font(FONT_B, 36)
        d.text((x + 36, 520), f"0{i+1}", font=nf, fill=ACCENT)
        d.text((x + 36, 600), s, font=sf, fill=TEXT)
        if i < 3:
            d.polygon([(x + 410, 570), (x + 440, 585), (x + 410, 600)], fill=ACCENT)


def draw_modes(d):
    kicker(d, "技术亮点")
    title(d, "三通道融合：有接口用接口，没有也能控", y=120, size=48)
    modes = [
        ("Mode-A  语义", "读取 AT-SPI 控件树，按角色和名称点击、输入。稳定，抗分辨率变化。", ACCENT),
        ("Mode-B  协议", "走 D-Bus、系统 Intent 和命令行。最精确，适合系统和已适配应用。", BLUE),
        ("Mode-C  合成", "窗口截图、中文识别、窗口内注入，文本经输入法提交。这是无接口软件的兜底。", AMBER),
    ]
    for i, (h, t, c) in enumerate(modes):
        card(d, 72, 340 + i * 210, 1760, 190, h, t, c)


def draw_agent(d):
    kicker(d, "产品亮点")
    title(d, "系统自带大脑，而不是外挂一个聊天框", y=130, size=50)
    cards = [
        ("本地优先", "隐私和离线任务走本机小模型，硬件不够就自动降级。"),
        ("先看计划", "复杂任务先给出可编辑步骤，批准后再动手。"),
        ("技能沉淀", "跑通的流程可以保存、参数化、下次一键重跑。"),
        ("可接云端", "长推理按数据分级路由，屏幕原图默认不上云。"),
    ]
    for i, (h, t) in enumerate(cards):
        x = 72 + (i % 2) * 900
        y = 380 + (i // 2) * 280
        card(d, x, y, 860, 250, h, t)


def draw_access(d):
    kicker(d, "接入亮点")
    title(d, "零改造可被控，也能接入你自己的智能体", y=130, size=48)
    rows = [
        ("L0  零改造", "靠系统 DCP 直接操控，应用不用改。"),
        ("MCP / CLI / SDK", "Cursor 等外部智能体调用同一套本机工具。"),
        ("自带模型", "可接 Ollama 或云端密钥，隐私分级仍由系统管。"),
        ("执行留在本机", "推理可以在外面，点击和截图策略不离开这台机器。"),
    ]
    for i, (h, t) in enumerate(rows):
        y = 360 + i * 150
        rounded(d, (72, y, 1848, y + 130), 16, CARD, LINE)
        d.text((110, y + 28), h, font=font(FONT_B, 32), fill=ACCENT)
        d.text((110, y + 76), t, font=font(FONT_R, 28), fill=MUTED)


def draw_rcp(d):
    kicker(d, "技术亮点")
    title(d, "网络和串口，同一套报文", y=130, size=56)
    body(d, "远程控制协议 RCP：TCP 与串口只是两条管道，命令和帧格式完全相同。", 280, 32)
    cols = [
        ("控桌面应用", "启动、列窗口、快照、查找、点击和输入。没有 AI 接口也走 DCP。"),
        ("控系统", "音量、锁屏、电源等白名单能力。没有任意命令行。"),
        ("控智能体", "下发任务、接收进度、随时中止。高风险动作仍要确认。"),
    ]
    for i, (h, t) in enumerate(cols):
        card(d, 72 + i * 600, 460, 560, 400, h, t, BLUE if i == 1 else ACCENT)


def draw_security(d):
    kicker(d, "安全亮点")
    title(d, "控制权在系统里，但不是后门", y=130, size=56)
    levels = [
        ("L0", "读公开状态", "会话内授权"),
        ("L1", "输入与切换", "任务级信任"),
        ("L2", "截图与写文件", "显式确认"),
        ("L3", "发送、删除、重启", "单次确认"),
    ]
    for i, (lv, name, rule) in enumerate(levels):
        x = 72 + i * 450
        rounded(d, (x, 400, x + 420, 760), 18, CARD, LINE)
        d.text((x + 36, 440), lv, font=font(FONT_B, 48), fill=ACCENT)
        d.text((x + 36, 530), name, font=font(FONT_B, 32), fill=TEXT)
        d.text((x + 36, 600), rule, font=font(FONT_R, 28), fill=MUTED)
    body(d, "密码、支付和系统认证界面默认拒绝自动化。全程审计，可一键停住。", 820, 30, TEXT)


def draw_feasible(d):
    kicker(d, "可行性")
    title(d, "站在 Linux 已有能力上，而不是空想", y=130, size=52)
    pts = [
        ("AT-SPI", "读屏软件已经用了几十年的控件树"),
        ("Wayland / Portal", "窗口、截图和输入都有系统级边界"),
        ("输入法", "中文提交走系统输入法，不靠模拟英文键"),
        ("隔离与审计", "cgroup、沙箱和策略，模型不进内核"),
    ]
    for i, (h, t) in enumerate(pts):
        y = 360 + i * 150
        d.ellipse((80, y + 18, 108, y + 46), fill=ACCENT)
        d.text((140, y), h, font=font(FONT_B, 34), fill=TEXT)
        d.text((140, y + 52), t, font=font(FONT_R, 28), fill=MUTED)


def draw_end(d):
    kicker(d, "AIOS")
    f = font(FONT_B, 58)
    lines = ["不在应用里找 AI 接口", "在操作系统里", "提供桌面操控平面"]
    y = 280
    for i, line in enumerate(lines):
        color = ACCENT if i == 2 else TEXT
        d.text((72, y), line, font=f, fill=color)
        y += 100
    body(d, "Linux 内核  ·  系统级 DCP  ·  自带 Agent  ·  网络与串口 RCP", 720, 32, MUTED)


DRAW = {
    "cover": draw_cover,
    "problem": draw_problem,
    "dcp": draw_dcp,
    "modes": draw_modes,
    "agent": draw_agent,
    "access": draw_access,
    "rcp": draw_rcp,
    "security": draw_security,
    "feasible": draw_feasible,
    "end": draw_end,
}


def render_slides():
    paths = []
    n = len(SLIDE_DEFS)
    for i, spec in enumerate(SLIDE_DEFS, 1):
        im, d = base()
        DRAW[spec["draw"]](d)
        footer(d, i, n)
        path = SLIDES / f"slide_{spec['id']}.png"
        im.save(path, "PNG")
        paths.append(path)
    return paths


def build_pptx(paths):
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for path, spec in zip(paths, SLIDE_DEFS):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(path), Emu(0), Emu(0), prs.slide_width, prs.slide_height)
        notes = slide.notes_slide.notes_text_frame
        notes.text = spec["narration"]
    out = ROOT / "AIOS宣传.pptx"
    prs.save(out)
    return out


async def synth_audio():
    files = []
    for spec in SLIDE_DEFS:
        dest = AUDIO / f"{spec['id']}.mp3"
        comm = edge_tts.Communicate(spec["narration"], VOICE, rate="-5%")
        await comm.save(str(dest))
        files.append(dest)
    return files


def duration(path):
    out = subprocess.check_output(
        ["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0", str(path)],
        text=True,
    )
    return float(out.strip())


def build_video(images, audios):
    # One clip per slide: still image + narration, short fade.
    clips = []
    for img, aud in zip(images, audios):
        dur = duration(aud) + 0.6
        clip = AUDIO / f"{img.stem}.mp4"
        fade_out = max(dur - 0.35, 0.1)
        subprocess.check_call(
            [
                "ffmpeg", "-y",
                "-loop", "1", "-i", str(img),
                "-i", str(aud),
                "-t", f"{dur:.3f}",
                "-vf", f"fade=t=in:st=0:d=0.25,fade=t=out:st={fade_out:.3f}:d=0.3,format=yuv420p",
                "-c:v", "libx264", "-pix_fmt", "yuv420p", "-r", "30",
                "-c:a", "aac", "-b:a", "192k",
                "-shortest",
                str(clip),
            ],
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
        clips.append(clip)
    lst = AUDIO / "concat.txt"
    lst.write_text("".join(f"file '{c.as_posix()}'\n" for c in clips), encoding="utf-8")
    out = ROOT / "AIOS宣传.mp4"
    subprocess.check_call(
        ["ffmpeg", "-y", "-f", "concat", "-safe", "0", "-i", str(lst), "-c", "copy", str(out)],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    return out


def main():
    paths = render_slides()
    ppt = build_pptx(paths)
    audios = asyncio.run(synth_audio())
    video = build_video(paths, audios)
    print(ppt)
    print(video)


if __name__ == "__main__":
    main()

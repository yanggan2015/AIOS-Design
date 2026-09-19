# -*- coding: utf-8 -*-
"""luminOS portrait promo: pain, solution, framework, new capabilities."""
import asyncio
import subprocess
from pathlib import Path

import edge_tts
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu, Inches

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "竖屏"
SLIDES = OUT / "slides"
AUDIO = OUT / "audio"
CLIPS = OUT / "clips"
for p in (SLIDES, AUDIO, CLIPS):
    p.mkdir(parents=True, exist_ok=True)

W, H = 1080, 1920
MX = 72
CW = W - MX * 2
BG_TOP = (24, 25, 28)
BG_BOT = (12, 12, 14)
CARD = (28, 29, 33)
LINE = (58, 54, 46)
GOLD = (214, 178, 112)
TEXT = (245, 242, 234)
MUTED = (166, 160, 150)
SOFT = (196, 190, 178)
FONT_B = "C:/Windows/Fonts/msyhbd.ttc"
FONT_R = "C:/Windows/Fonts/msyh.ttc"
VOICE = "zh-CN-YunxiNeural"

SLIDES_DATA = [
    {
        "narration": "这是 luminOS。一套专门为人工智能体设计的操作系统。它要解决的不是再做一个聊天窗口，而是让智能体真正能够使用这台电脑上的软件。",
        "draw": "cover",
    },
    {
        "narration": "你现在用的智能体，已经很会想、很会写、很会做计划。可一到真正的桌面软件，它常常只能停在外面。不是它不够聪明，是这台电脑还没有给它一双手。",
        "draw": "scene",
    },
    {
        "narration": "痛点就三件。软件不提供人工智能接口，智能体进不去。每个应用单独做插件，永远做不完。靠坐标去点的脚本，界面一改就失效，也没有确认，没有审计。",
        "draw": "pains",
    },
    {
        "narration": "luminOS 不要求每一个软件先为智能体开门。操作能力由操作系统提供。应用不用改造，安装上来，就可以被操控。",
        "draw": "turn",
    },
    {
        "narration": "三个痛点，是这样解开的。进不去，就用语义、协议和视觉三条路，没有接口也能操作窗口。接不完，是因为能力做在系统里，做一次，所有应用都能用。脚本又脆又不安全，我们就让每一步都有目标、有检查、有确认，事后也能查到。",
        "draw": "solve",
    },
    {
        "narration": "框架可以记成四层。你提出目标。智能体负责理解和计划。桌面操控平面负责找到窗口并动手。最下面仍然是普通的 Linux 应用。人工智能引擎负责思考，不放进内核。",
        "draw": "frame",
    },
    {
        "narration": "操控平面里，能读到控件，就按控件操作。系统本来就有协议，就走协议。两者都没有，就只看这一个窗口，识别文字，再在窗口里面点击和输入。中文通过系统输入法送进去。",
        "draw": "channels",
    },
    {
        "narration": "在这个框架上，有三件现在特别需要的能力。网络和串口使用同一套报文，可以远程控制应用和系统。你可以只用系统自带的智能体，也可以接入你已经在用的智能体，真正的执行仍然留在本机。高风险动作必须确认。密码和支付界面，默认不做自动化。",
        "draw": "new",
    },
    {
        "narration": "luminOS。为人工智能体设计的操作系统。不在应用里找接口。在操作系统里，给智能体一双手。",
        "draw": "end",
    },
]


def font(path, size):
    return ImageFont.truetype(path, size)


def wrap(draw, text, fnt, max_w):
    lines, cur = [], ""
    for ch in text:
        trial = cur + ch
        if draw.textlength(trial, font=fnt) <= max_w:
            cur = trial
        else:
            if cur:
                lines.append(cur)
            cur = ch
    if cur:
        lines.append(cur)
    return lines


def draw_lines(d, lines, x, y, fnt, fill, size, leading):
    for i, line in enumerate(lines):
        d.text((x, y + i * (size + leading)), line, font=fnt, fill=fill)
    return y + len(lines) * (size + leading)


def base():
    col = Image.new("RGB", (1, H))
    px = col.load()
    for y in range(H):
        t = y / (H - 1)
        px[0, y] = tuple(int(BG_TOP[i] + (BG_BOT[i] - BG_TOP[i]) * t) for i in range(3))
    im = col.resize((W, H), Image.Resampling.BILINEAR)
    d = ImageDraw.Draw(im)
    d.rectangle((MX, 64, MX + 48, 68), fill=GOLD)
    return im, d


def footer(d, n, total):
    d.text((MX, 1844), "luminOS", font=font(FONT_B, 22), fill=GOLD)
    d.text((MX + 132, 1846), "为 AI Agent 设计", font=font(FONT_R, 22), fill=MUTED)
    label = f"{n:02d}  /  {total:02d}"
    f = font(FONT_R, 22)
    d.text((W - MX - d.textlength(label, font=f), 1846), label, font=f, fill=MUTED)


def kicker(d, text, y=112):
    d.text((MX, y), text, font=font(FONT_B, 26), fill=GOLD)
    return y + 56


def heading(d, text, y, size=52):
    f = font(FONT_B, size)
    lines = wrap(d, text, f, CW)
    return draw_lines(d, lines, MX, y, f, TEXT, size, 14)


def paragraph(d, text, y, size=30, color=MUTED, width=None, leading=16):
    f = font(FONT_R, size)
    lines = wrap(d, text, f, width or CW)
    return draw_lines(d, lines, MX, y, f, color, size, leading)


def card(d, x, y, w, h):
    d.rounded_rectangle((x, y, x + w, y + h), radius=20, fill=CARD, outline=LINE, width=2)


def draw_cover(d):
    kicker(d, "一套 AI 操作系统", 220)
    d.text((MX, 340), "luminOS", font=font(FONT_B, 92), fill=TEXT)
    d.rectangle((MX, 470, MX + 88, 474), fill=GOLD)
    y = heading(d, "为 AI Agent 设计", 520, 56)
    paragraph(d, "让智能体不再停在软件外面，而是能够使用这台电脑。", y + 36, 34, SOFT, leading=18)


def draw_scene(d):
    y = kicker(d, "现在的痛点")
    y = heading(d, "智能体已经很会想", y + 24, 56)
    y = heading(d, "一到软件，就停住", y + 8, 56)
    d.rectangle((MX, y + 48, MX + 88, y + 52), fill=GOLD)
    paragraph(d, "不是它不够聪明。是操作系统还没有给它一双手。", y + 96, 34, SOFT, leading=18)


def draw_pains(d):
    y = kicker(d, "三个具体问题")
    y = heading(d, "为什么今天控不了桌面", y + 8, 48)
    items = [
        ("01", "进不去", "软件不提供 AI 接口，智能体只能停在窗口外面。"),
        ("02", "接不完", "每个应用单独做插件，覆盖速度取决于厂商。"),
        ("03", "又脆又不安全", "坐标脚本一改版就失效，没有确认，也没有审计。"),
    ]
    y += 36
    for num, title, text in items:
        card(d, MX, y, CW, 360)
        d.text((MX + 36, y + 36), num, font=font(FONT_B, 26), fill=GOLD)
        d.text((MX + 36, y + 88), title, font=font(FONT_B, 42), fill=TEXT)
        bf = font(FONT_R, 30)
        lines = wrap(d, text, bf, CW - 72)
        draw_lines(d, lines, MX + 36, y + 170, bf, MUTED, 30, 14)
        y += 384


def draw_turn(d):
    y = kicker(d, "我们的做法", 200)
    lines = ["不要求软件", "为智能体开门"]
    f = font(FONT_B, 64)
    yy = 340
    for line in lines:
        d.text((MX, yy), line, font=f, fill=TEXT)
        yy += 96
    d.rectangle((MX, yy + 20, MX + 88, yy + 24), fill=GOLD)
    paragraph(d, "操作能力由操作系统提供。应用零改造，安装上来就可以被操控。", yy + 72, 34, SOFT, leading=18)


def draw_solve(d):
    y = kicker(d, "痛点如何解开")
    y = heading(d, "一个痛点，对应一个做法", y + 4, 46)
    pairs = [
        ("进不去", "语义、协议、视觉三条路。没有接口，也能操作窗口。"),
        ("接不完", "能力做在系统里。做一次，所有应用都能用。"),
        ("又脆又不安全", "每一步有目标、有检查、有确认，事后也能查到。"),
    ]
    y += 28
    for pain, fix in pairs:
        card(d, MX, y, CW, 390)
        d.text((MX + 36, y + 32), "痛点", font=font(FONT_R, 24), fill=MUTED)
        d.text((MX + 36, y + 72), pain, font=font(FONT_B, 40), fill=TEXT)
        d.text((MX + 36, y + 160), "做法", font=font(FONT_R, 24), fill=GOLD)
        bf = font(FONT_R, 30)
        lines = wrap(d, fix, bf, CW - 72)
        draw_lines(d, lines, MX + 36, y + 208, bf, SOFT, 30, 14)
        y += 414


def draw_frame(d):
    y = kicker(d, "解决框架")
    y = heading(d, "四层就够记住", y + 4, 52)
    layers = [
        ("01", "你", "说出要完成的事"),
        ("02", "智能体", "理解、计划、决定下一步"),
        ("03", "操控平面", "找到窗口，并真正动手"),
        ("04", "桌面应用", "不用改造，普通 Linux 软件"),
    ]
    y += 28
    for i, (num, name, desc) in enumerate(layers):
        card(d, MX, y, CW, 230)
        d.text((MX + 36, y + 78), num, font=font(FONT_B, 28), fill=GOLD)
        d.text((MX + 140, y + 48), name, font=font(FONT_B, 40), fill=TEXT)
        d.text((MX + 140, y + 120), desc, font=font(FONT_R, 28), fill=MUTED)
        y += 250
    paragraph(d, "人工智能引擎负责思考，不放进内核。", y + 8, 28, SOFT)


def draw_channels(d):
    y = kicker(d, "框架里的关键机制")
    y = heading(d, "有接口就用，没有也能控", y + 4, 48)
    modes = [
        ("先", "读控件", "界面自己暴露了按钮和文字，就按名称操作。稳，也不怕分辨率变化。"),
        ("再", "走协议", "系统本来就有的消息和命令，优先精确调用。"),
        ("最后", "看窗口", "截取这一个窗口，识别文字，在窗口内点击和输入。中文走输入法。"),
    ]
    y += 28
    for mark, name, text in modes:
        card(d, MX, y, CW, 400)
        d.text((MX + 36, y + 32), mark, font=font(FONT_B, 26), fill=GOLD)
        d.text((MX + 36, y + 84), name, font=font(FONT_B, 44), fill=TEXT)
        bf = font(FONT_R, 30)
        lines = wrap(d, text, bf, CW - 72)
        draw_lines(d, lines, MX + 36, y + 170, bf, MUTED, 30, 14)
        y += 424


def draw_new(d):
    y = kicker(d, "在此之上的新能力")
    y = heading(d, "远程、接入、边界", y + 4, 52)
    items = [
        ("同一套报文", "网络和串口命令相同。可以远程控制应用，也可以控制系统。不是第二套后门。"),
        ("智能体可替换", "用系统自带的，或接入你已经在用的。点击和权限仍留在本机。"),
        ("危险动作要确认", "发送、删除、重启必须确认。密码和支付界面默认不自动操作。"),
    ]
    y += 28
    for title, text in items:
        card(d, MX, y, CW, 400)
        d.text((MX + 36, y + 36), title, font=font(FONT_B, 40), fill=TEXT)
        bf = font(FONT_R, 30)
        lines = wrap(d, text, bf, CW - 72)
        draw_lines(d, lines, MX + 36, y + 120, bf, MUTED, 30, 14)
        y += 424


def draw_end(d):
    kicker(d, "luminOS", 300)
    f = font(FONT_B, 52)
    lines = [("不在应用里找接口", TEXT), ("在操作系统里", TEXT), ("给智能体一双手", GOLD)]
    y = 460
    for line, color in lines:
        wrapped = wrap(d, line, f, CW)
        y = draw_lines(d, wrapped, MX, y, f, color, 52, 16)
        y += 12
    d.rectangle((MX, y + 28, MX + 88, y + 32), fill=GOLD)
    paragraph(d, "为 AI Agent 设计的 AI 操作系统", y + 72, 32, SOFT)


DRAW = {
    "cover": draw_cover,
    "scene": draw_scene,
    "pains": draw_pains,
    "turn": draw_turn,
    "solve": draw_solve,
    "frame": draw_frame,
    "channels": draw_channels,
    "new": draw_new,
    "end": draw_end,
}


def render():
    paths = []
    total = len(SLIDES_DATA)
    for i, spec in enumerate(SLIDES_DATA, 1):
        im, d = base()
        DRAW[spec["draw"]](d)
        footer(d, i, total)
        path = SLIDES / f"slide_{i:02d}.png"
        im.save(path, "PNG")
        paths.append(path)
    return paths


def build_pptx(paths):
    prs = Presentation()
    prs.slide_width = Inches(7.5)
    prs.slide_height = Inches(13.333333)
    blank = prs.slide_layouts[6]
    for path, spec in zip(paths, SLIDES_DATA):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(path), Emu(0), Emu(0), prs.slide_width, prs.slide_height)
        slide.notes_slide.notes_text_frame.text = spec["narration"]
    out = OUT / "luminOS宣传-竖屏.pptx"
    prs.save(out)
    return out


async def synth():
    files = []
    for i, spec in enumerate(SLIDES_DATA, 1):
        dest = AUDIO / f"{i:02d}.mp3"
        await edge_tts.Communicate(spec["narration"], VOICE, rate="-6%").save(str(dest))
        files.append(dest)
    return files


def duration(path):
    out = subprocess.check_output(
        ["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0", str(path)],
        text=True,
    )
    return float(out.strip())


def build_video(images, audios):
    clips = []
    for img, aud in zip(images, audios):
        dur = duration(aud) + 0.45
        clip = CLIPS / f"{img.stem}.mp4"
        fade_out = max(dur - 0.28, 0.1)
        subprocess.check_call(
            [
                "ffmpeg", "-y", "-loop", "1", "-i", str(img), "-i", str(aud),
                "-t", f"{dur:.3f}",
                "-vf", f"scale=1080:1920,fade=t=in:st=0:d=0.25,fade=t=out:st={fade_out:.3f}:d=0.25,format=yuv420p",
                "-c:v", "libx264", "-pix_fmt", "yuv420p", "-r", "30",
                "-c:a", "aac", "-b:a", "192k", "-shortest", str(clip),
            ],
            stdout=subprocess.DEVNULL,
        )
        clips.append(clip)
    lst = CLIPS / "concat.txt"
    lst.write_text("".join(f"file '{c.as_posix()}'\n" for c in clips), encoding="utf-8")
    video = OUT / "luminOS宣传-竖屏.mp4"
    subprocess.check_call(
        ["ffmpeg", "-y", "-f", "concat", "-safe", "0", "-i", str(lst), "-c", "copy", str(video)],
        stdout=subprocess.DEVNULL,
    )
    return video


def main():
    paths = render()
    ppt = build_pptx(paths)
    audios = asyncio.run(synth())
    video = build_video(paths, audios)
    print(ppt)
    print(video)


if __name__ == "__main__":
    main()
